"""Chuyển chứng thư từ PDF sang DOCX, giữ nguyên nội dung.

Chạy:
    ./.venv/bin/python3 -m fixtures.convert_certificate_pdf_to_docx \
        "samples/CT-SAMPLE-001 - CHUNG-THU-TDG-ANONYMIZED.pdf" samples/docx

Vì sao dựng lại thay vì "convert" bằng công cụ ngoài: LibreOffice nhập PDF qua bộ
lọc vẽ hình, cho ra DOCX gồm các hộp text rời rạc — mở lên trông giống nhưng
không còn đoạn văn và bảng THẬT, nên dùng làm input thử nghiệm thì vô nghĩa.
Pandoc thì không đọc được PDF.

Cách ở đây: lấy chính cấu trúc mà pipeline đã bóc (mục La Mã, cặp nhãn-giá trị,
bảng theo đường kẻ ô) rồi dựng lại thành Word với đoạn văn và bảng thật. Nội dung
giữ nguyên từng ký tự, còn cấu trúc thì đúng nghĩa Word — đúng thứ cần để thử
nhánh xử lý DOCX.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from pdf_extract.pipeline import process_document

from .docx_verbatim_lookups import (
    bullet_markers,
    heading_lines,
    labels_with_colon,
    marker_for,
)

# Vị trí tab để cột giá trị thẳng hàng, xấp xỉ bố cục của chứng thư gốc.
VALUE_TAB_INCHES = 1.75

# Cỡ chữ thân tài liệu, theo bản PDF.
BODY_FONT_SIZE_PT = 12
TITLE_FONT_SIZE_PT = 15


def convert(pdf_path: str, output_dir: Path) -> Path:
    """Dựng file DOCX từ nội dung đã bóc của một chứng thư PDF."""
    result = process_document(pdf_path)
    if not result.data:
        raise ValueError(f"Không bóc được nội dung từ {pdf_path}: {result.errors}")

    document = Document()
    _apply_base_style(document)

    data = result.data
    # Ký tự mở đầu mục liệt kê ("-" hay "+") bị parser tách ra khỏi giá trị, nên
    # tra lại từ text nguyên văn để bản DOCX giữ đúng dấu của tài liệu.
    markers = bullet_markers(data)
    headings = heading_lines(data)
    colon_labels = labels_with_colon(data)

    _write_preface(document, data["preface"], markers)

    for number, section in data["sections"].items():
        _write_section(
            document, number, section, markers, headings.get(number), colon_labels
        )

    _write_signatures(document, data["signatures"])
    _write_page_footer(document, data["page_footers"])

    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"{Path(pdf_path).stem}.docx"
    document.save(str(target))

    return target


def _write_page_footer(document: Document, footers: list[dict]) -> None:
    """Chân trang. Dùng field PAGE/NUMPAGES — cách Word đánh số trang.

    Bản PDF ghi cố định "Trang 1/3", "Trang 2/3", "Trang 3/3". DOCX không có
    trang cố định (Word phân trang khi mở file), nên số trang phải là FIELD.
    Kết quả hiển thị giống hệt, chỉ khác ở chỗ nó được tính lúc mở file.
    """
    if not footers:
        return

    label = footers[0]["value"].split()[0]  # "Trang"
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run(f"{label} ").bold = True

    _add_field(paragraph, "PAGE")
    paragraph.add_run("/")
    _add_field(paragraph, "NUMPAGES")


def _add_field(paragraph, instruction: str) -> None:
    """Chèn một field Word (PAGE, NUMPAGES) vào đoạn văn."""
    run = paragraph.add_run()

    begin = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    text = run._r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
    text.text = f" {instruction} "
    end = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})

    for element in (begin, text, end):
        run._r.append(element)


def _apply_base_style(document: Document) -> None:
    """Đặt font và cỡ chữ giống bản PDF (Times New Roman 12pt)."""
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(BODY_FONT_SIZE_PT)


def _paragraph(document: Document, text: str, *, bold: bool = False, align=None, size=None):
    """Một đoạn văn thường, tuỳ chọn in đậm và căn lề."""
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold

    if size is not None:
        run.font.size = Pt(size)
    if align is not None:
        paragraph.alignment = align

    return paragraph


def _label_value_paragraph(
    document: Document, label: str, value: str | None, marker: str = "-"
) -> None:
    """Dòng "nhãn<tab>: giá trị" — giữ đúng cách chứng thư căn hai cột."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(VALUE_TAB_INCHES * 72))

    paragraph.add_run(f"{marker} {label}".strip())
    paragraph.add_run("\t: ")

    if value:
        paragraph.add_run(value)


def _write_preface(document: Document, preface: dict, markers: dict[str, str]) -> None:
    """Tiêu đề thư, quốc hiệu, số hợp đồng, "Kính gửi", các dòng "Căn cứ"."""
    paragraphs = [entry["value"] for entry in preface["paragraphs"]]
    fields = preface["fields"]

    # Hai dòng đầu của chứng thư là hai cột: đơn vị bên trái, quốc hiệu bên phải.
    if paragraphs:
        _paragraph(document, paragraphs[0], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for text in paragraphs[1:3]:
        _paragraph(document, text, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    for name in ("contract_number", "certificate_number"):
        entry = fields.get(name)
        if entry and entry["value"]:
            _paragraph(document, f"{entry['label']}: {entry['value']}")

    for text in paragraphs[3:4]:
        _paragraph(document, text, align=WD_ALIGN_PARAGRAPH.RIGHT)

    for text in paragraphs[4:]:
        _paragraph(
            document,
            text,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            size=TITLE_FONT_SIZE_PT,
        )

    recipient = fields.get("recipient")
    if recipient and recipient["value"]:
        _paragraph(document, f"{recipient['label']}: {recipient['value']}", bold=True)

    for item in preface["items"]:
        _paragraph(document, f"{marker_for(markers, item['value'])} {item['value']}".strip())


def _write_section(
    document: Document,
    number: str,
    section: dict,
    markers: dict[str, str],
    heading_line: str | None,
    colon_labels: set[str],
) -> None:
    """Một mục: tiêu đề, giá trị của tiêu đề, các field, đoạn văn, mục liệt kê, bảng."""
    raw_heading = heading_line or f"{number}. {section['title']['value']}"
    prefix, _, rest = raw_heading.partition(":")
    inline = section.get("value")

    # Tiêu đề viết kèm giá trị ngay sau dấu hai chấm: dùng GIÁ TRỊ ĐẦY ĐỦ từ
    # JSON, không dùng phần còn lại của dòng tiêu đề — giá trị có thể vắt sang
    # dòng dưới nên dòng tiêu đề chỉ chứa đoạn đầu.
    if rest.strip() and inline and inline["value"]:
        _paragraph(document, f"{prefix}: {inline['value']}", bold=True)
    else:
        _paragraph(document, raw_heading, bold=True)

    for entry in {**section["fields"], **section["unmapped_fields"]}.values():
        _label_value_paragraph(
            document,
            entry["label"],
            entry.get("value"),
            marker_for(markers, entry["label"]),
        )

    for paragraph in section["paragraphs"]:
        _paragraph(document, paragraph["value"])

    if "table" in section:
        _write_asset_table(document, section["table"], colon_labels)

    for item in section["items"]:
        _paragraph(document, f"{marker_for(markers, item['value'])} {item['value']}".strip())


def _write_asset_table(
    document: Document, table_data: dict, colon_labels: set[str]
) -> None:
    """Bảng giá trị tài sản: tiêu đề cột, các thửa đất, các dòng tổng gộp ô."""
    headers = [column["value"] for column in table_data["columns"]]
    rows = table_data["rows"]

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell, header in zip(table.rows[0].cells, headers):
        cell.paragraphs[0].add_run(header).bold = True

    for row in rows:
        cells = table.add_row().cells
        values = [
            str(row["index"]),
            row["description"]["value"],
            row["area_sqm"]["value"],
            row["amount_vnd"]["value"],
        ]
        for cell, value in zip(cells, values):
            cell.text = value

    for entry in table_data["totals"].values():
        if entry is None:
            continue

        cells = table.add_row().cells
        # Gộp ô đúng như bản PDF: nhãn chiếm hai cột đầu, giá trị hai cột cuối.
        label_cell = cells[0].merge(cells[1])
        value_cell = cells[2].merge(cells[3])

        label = entry["label"]
        label_cell.text = f"{label}:" if label in colon_labels else label
        value_cell.text = entry.get("value") or ""


def _write_signatures(document: Document, signatures: list[dict]) -> None:
    """Khối chữ ký: mỗi người một cột, giữ đúng cách xếp cạnh nhau."""
    if not signatures:
        return

    document.add_paragraph()
    table = document.add_table(rows=1, cols=len(signatures))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell, signature in zip(table.rows[0].cells, signatures):
        first = True
        for line in signature["lines"]:
            paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(line["value"]).bold = True
            first = False


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 2

    output_dir = Path(argv[2]) if len(argv) > 2 else Path("samples/docx")
    target = convert(argv[1], output_dir)
    print(f"  {target}  ({target.stat().st_size:,} bytes)")

    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
