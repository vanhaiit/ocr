"""Đọc nội dung DOCX: đoạn văn, bảng, chân trang — kèm đường dẫn XML.

Khác PDF ở chỗ căn bản: PDF cho một mặt phẳng glyph có toạ độ, phải tự gom thành
dòng và tự dò biên bảng. DOCX thì Word đã ghi sẵn cấu trúc — mỗi `<w:p>` là một
đoạn HOÀN CHỈNH (Word tự ngắt dòng khi hiển thị, nên không có chuyện giá trị bị
cắt làm hai), và mỗi `<w:tbl>` là một bảng có ô rõ ràng.

Vì vậy DOCX **dễ bóc hơn** nhưng **khó chứng minh hơn**: không có toạ độ để làm
bằng chứng vị trí. Thay bằng ĐƯỜNG DẪN XML (`body/p[12]`,
`body/tbl[1]/tr[3]/tc[2]`) — nó chỉ đúng một node trong file, nên vẫn truy được
về nguồn, chỉ là theo cấu trúc thay vì theo hình học.

Đọc cả chân trang: nội dung chân trang là text của tài liệu, bỏ qua là mất dữ
liệu.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from docx import Document
from docx.oxml.ns import qn

from .models import SourcedValue

# DOCX không phân trang (Word tính lúc mở file), nên không có số trang thật.
# Dùng 0 để phân biệt rõ với trang 1 của PDF.
UNPAGINATED = 0


@dataclass
class DocxParagraph:
    """Một đoạn văn của tài liệu, kèm đường dẫn XML và mức thụt lề."""

    text: str
    location: str
    # Mức danh sách của Word (`w:ilvl`), None nếu không phải danh sách.
    list_level: int | None = None

    def as_value(self) -> SourcedValue:
        return SourcedValue(value=self.text, page=UNPAGINATED, location=self.location)


@dataclass
class DocxTable:
    """Một bảng: các hàng, mỗi ô là một giá trị có nguồn gốc."""

    location: str
    rows: list[list[SourcedValue]] = field(default_factory=list)

    @property
    def shape(self) -> tuple[int, int]:
        return (len(self.rows), max((len(r) for r in self.rows), default=0))


@dataclass
class DocxContent:
    """Toàn bộ nội dung đọc được từ một DOCX."""

    paragraphs: list[DocxParagraph] = field(default_factory=list)
    tables: list[DocxTable] = field(default_factory=list)
    footers: list[SourcedValue] = field(default_factory=list)
    # Thứ tự xuất hiện của đoạn văn và bảng trong thân tài liệu. Cần để parser
    # biết bảng nằm ở mục nào, thay vì phải đoán.
    order: list[tuple[str, int]] = field(default_factory=list)
    # Trường động Word (PAGE, NUMPAGES, DATE...) không có giá trị đã tính sẵn
    # trong file — Word chỉ tính khi mở file, XML không có gì để đọc. Ghi lại
    # để báo lỗi thay vì âm thầm trả về text thiếu ký tự nhưng vẫn coi là đủ.
    unresolved_fields: list[str] = field(default_factory=list)


def extract_docx(path: str) -> DocxContent:
    """Đọc đoạn văn, bảng và chân trang của một file DOCX."""
    document = Document(path)
    content = DocxContent()

    _read_body(document, content)
    _read_footers(document, content)

    return content


def _read_body(document: Document, content: DocxContent) -> None:
    """Duyệt thân tài liệu theo ĐÚNG thứ tự đoạn văn và bảng xuất hiện."""
    paragraph_index = 0
    table_index = 0

    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph_index += 1
            text = _element_text(element)
            location = f"body/p[{paragraph_index}]"
            _record_unresolved_fields(element, location, content)

            if not text.strip():
                continue

            content.paragraphs.append(
                DocxParagraph(
                    text=text,
                    location=location,
                    list_level=_list_level(element),
                )
            )
            content.order.append(("paragraph", len(content.paragraphs) - 1))
            continue

        if element.tag == qn("w:tbl"):
            table_index += 1
            content.tables.append(
                _read_table(element, f"body/tbl[{table_index}]", content)
            )
            content.order.append(("table", len(content.tables) - 1))


def _read_table(element, location: str, content: DocxContent) -> DocxTable:
    """Đọc một bảng; ô gộp chỉ lấy MỘT lần dù nó trải trên nhiều cột."""
    table = DocxTable(location=location)

    for row_index, row in enumerate(element.iterchildren(qn("w:tr")), start=1):
        cells: list[SourcedValue] = []

        for cell_index, cell in enumerate(row.iterchildren(qn("w:tc")), start=1):
            cell_location = f"{location}/tr[{row_index}]/tc[{cell_index}]"
            _record_unresolved_fields(cell, cell_location, content)
            cells.append(
                SourcedValue(
                    value=_element_text(cell),
                    page=UNPAGINATED,
                    location=cell_location,
                )
            )

        table.rows.append(cells)

    return table


def _read_footers(document: Document, content: DocxContent) -> None:
    """Chân trang của từng section. Là text của tài liệu nên phải lấy."""
    for index, section in enumerate(document.sections, start=1):
        footer_element = section.footer._element
        location = f"sectPr[{index}]/footer"
        _record_unresolved_fields(footer_element, location, content)

        text = _element_text(footer_element).strip()
        if not text:
            continue

        content.footers.append(
            SourcedValue(
                value=text, page=UNPAGINATED, location=f"sectPr[{index}]/footer"
            )
        )


def _element_text(element) -> str:
    """Ghép mọi `<w:t>` bên dưới một node, giữ ranh giới đoạn và tab.

    Ba loại ranh giới phải giữ:

      - `<w:tab>`  -> khoảng trắng. Tab là cách Word căn cột cho cặp nhãn-giá
        trị; đổi thành khoảng trắng để tầng phân tích chỉ xử lý một loại phân
        cách.
      - `<w:br>`   -> xuống dòng trong cùng một đoạn.
      - `<w:p>`    -> xuống dòng GIỮA các đoạn. Cần thiết vì ô bảng chứa nhiều
        đoạn: ô chữ ký có vai trò, số thẻ, "(Ký tên)", họ tên mỗi thứ một đoạn,
        nối liền không ngắt sẽ thành một chuỗi vô nghĩa.
    """
    parts: list[str] = []
    paragraph_tag = qn("w:p")

    for node in element.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append(" ")
        elif node.tag == qn("w:br"):
            parts.append("\n")
        elif node.tag == paragraph_tag and node is not element and parts:
            parts.append("\n")

    return "".join(parts)


def _record_unresolved_fields(element, location: str, content: DocxContent) -> None:
    """Ghi lại field code (PAGE, NUMPAGES, DATE...) không có giá trị đã tính sẵn.

    Field phức hợp của Word là bốn phần theo thứ tự: `fldChar[begin]`,
    `instrText` (mã lệnh, ví dụ `PAGE`), `fldChar[separate]`, rồi `<w:t>` chứa
    KẾT QUẢ đã tính — Word ghi kết quả này khi lưu file để hiển thị ngay lần mở
    sau mà không cần tính lại. Nếu file chưa từng được Word tính (sinh bằng
    python-docx, chưa mở qua Word) thì không có phần kết quả, `_element_text`
    đọc field code đó ra chuỗi rỗng — mất ký tự mà không có gì báo hiệu.
    """
    fld_char = qn("w:fldChar")
    instr_text = qn("w:instrText")
    text_tag = qn("w:t")

    in_field = False
    has_result = False
    code = ""

    for node in element.iter():
        if node.tag == fld_char:
            kind = node.get(qn("w:fldCharType"))
            if kind == "begin":
                in_field, has_result, code = True, False, ""
            elif kind == "end" and in_field:
                if not has_result:
                    content.unresolved_fields.append(f"{location}: {code.strip() or '?'}")
                in_field = False
        elif in_field and node.tag == instr_text:
            code += node.text or ""
        elif in_field and node.tag == text_tag and node.text:
            has_result = True


def _list_level(element) -> int | None:
    """Mức danh sách của Word, để nhận mục liệt kê không có ký tự gạch đầu dòng."""
    level = element.find(f".//{qn('w:numPr')}/{qn('w:ilvl')}")
    if level is None:
        return None

    raw = level.get(qn("w:val"))
    return int(raw) if raw is not None and raw.isdigit() else 0


def canonical_docx_text(content: DocxContent) -> str:
    """Text chuẩn của tài liệu, dùng làm căn cứ cho cổng nguồn gốc.

    Ghép theo đúng thứ tự xuất hiện, rồi tới chân trang. Ô bảng mỗi ô một dòng
    để giá trị trong ô vẫn là chuỗi liền mạch khi cổng đi tìm.
    """
    lines: list[str] = []

    for kind, index in content.order:
        if kind == "paragraph":
            lines.append(content.paragraphs[index].text)
            continue

        for row in content.tables[index].rows:
            lines.extend(cell.value for cell in row if cell.value.strip())

    lines.extend(footer.value for footer in content.footers)

    return "\n".join(lines)
